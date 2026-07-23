from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def history_signature(payload: dict) -> dict[str, tuple[str, tuple[str, ...]]]:
    result: dict[str, tuple[str, tuple[str, ...]]] = {}
    for order_number, record in (payload.get("orders") or {}).items():
        version_ids = tuple(
            str(version.get("version_id") or version.get("id") or "")
            for version in (record.get("versions") or [])
        )
        result[str(order_number)] = (str(record.get("current_version_id") or ""), version_ids)
    return result


def workbook_values(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        return workbook.sheetnames, rows
    finally:
        workbook.close()


def csv_rows(path: Path) -> int:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return max(0, sum(1 for _ in csv.reader(stream)) - 1)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-dir", required=True)
    parser.add_argument("--run-dir", required=True)
    parser.add_argument("--backup-dir", required=True)
    args = parser.parse_args()

    state_dir = Path(args.state_dir)
    run_dir = Path(args.run_dir)
    backup_dir = Path(args.backup_dir)
    comparison = read_json(run_dir / "comparison_preview.json")
    before_history = read_json(backup_dir / "order_history.json")
    after_history = read_json(state_dir / "order_history.json")
    before_signature = history_signature(before_history)
    after_signature = history_signature(after_history)
    before_sheets, before_rows = workbook_values(backup_dir / "current_orders.xlsx")
    after_sheets, after_rows = workbook_values(state_dir / "current_orders.xlsx")

    checks = {
        "comparison_new_zero": int(comparison["summary"].get("new_count") or 0) == 0,
        "comparison_changed_zero": int(comparison["summary"].get("changed_count") or 0) == 0,
        "history_order_ids_unchanged": before_signature.keys() == after_signature.keys(),
        "history_version_ids_unchanged": before_signature == after_signature,
        "change_log_unchanged": sha256(backup_dir / "change_log.csv") == sha256(state_dir / "change_log.csv"),
        "excel_sheet_names_unchanged": before_sheets == after_sheets,
        "excel_rows_unchanged": before_rows == after_rows,
    }
    result = {
        "status": "passed" if all(checks.values()) else "failed",
        "checks": checks,
        "before_history_orders": len(before_signature),
        "after_history_orders": len(after_signature),
        "before_history_versions": sum(len(value[1]) for value in before_signature.values()),
        "after_history_versions": sum(len(value[1]) for value in after_signature.values()),
        "change_log_rows": csv_rows(state_dir / "change_log.csv"),
        "excel_sheet_names": after_sheets,
        "excel_data_rows": max(0, len(after_rows) - 1),
        "backup_dir": str(backup_dir.resolve()),
        "state_dir": str(state_dir.resolve()),
    }
    output = run_dir / "same_day_replay_verification.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
