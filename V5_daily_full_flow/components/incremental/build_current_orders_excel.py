from __future__ import annotations

import json
import os
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


DETAIL_HEADERS = [
    "門市",
    "訂單分類",
    "訂單編號",
    "日期",
    "品名",
    "定價",
    "數量",
    "總金額",
    "折後",
]
SUMMARY_CATEGORIES = ["婚約", "彌月", "企業", "其他"]
SUMMARY_HEADERS = [
    "訂單分類",
    "訂單數",
    "禮盒總數",
    "總金額",
    "平均訂單金額",
    "平均單盒金額",
    "平均訂單盒數",
]
CATEGORY_COLORS = {
    "婚約": "FFF2CC",
    "彌月": "E2F0D9",
    "企業": "F4CCCC",
    "其他": "D9EAF7",
    "Total": "FFFF00",
}
STORE_BY_DEPARTMENT = {"4001": "SR"}

REFERENCE_DIR = Path(__file__).resolve().parents[2] / "reference_data"
PRODUCT_LOOKUP_PATH = REFERENCE_DIR / "product_lookup.json"
MENUITEM_PATH = REFERENCE_DIR / "MENUITEM.XLS"
sys.path.insert(0, str(REFERENCE_DIR))
from menuitem_lookup import load_or_refresh


def parse_date(value: object) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def display_category(value: object) -> str:
    category = str(value or "").strip()
    if category == "節慶":
        return "企業"
    if category in SUMMARY_CATEGORIES:
        return category
    return "其他"


def parse_number(value: object) -> Decimal:
    text = str(value or "0").replace(",", "").strip()
    try:
        return Decimal(text or "0")
    except InvalidOperation:
        return Decimal(0)


def round_half_up(numerator: Decimal, denominator: int | Decimal) -> int:
    divisor = Decimal(str(denominator or 0))
    if not divisor:
        return 0
    return int((numerator / divisor).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def report_as_of() -> datetime:
    configured = os.environ.get("V5_REPORT_AS_OF", "").strip()
    if configured:
        parsed = parse_date(configured)
        if parsed:
            return parsed
    return datetime.now()


def recent_month_keys(as_of: datetime, count: int = 6) -> list[str]:
    keys = []
    for offset in range(count - 1, -1, -1):
        absolute_month = as_of.year * 12 + as_of.month - 1 - offset
        year, month_index = divmod(absolute_month, 12)
        keys.append(f"{year:04d}-{month_index + 1:02d}")
    return keys


def empty_metrics() -> dict[str, object]:
    return {"order_numbers": set(), "boxes": Decimal(0), "amount": Decimal(0)}


def calculate_month_summary(orders: list[dict]) -> dict[str, dict[str, object]]:
    summary = {category: empty_metrics() for category in SUMMARY_CATEGORIES}
    for order in orders:
        category = display_category(order.get("category"))
        metrics = summary[category]
        order_number = str(order.get("order_no") or "").strip()
        if order_number:
            metrics["order_numbers"].add(order_number)
        for detail in order.get("details") or []:
            metrics["boxes"] += parse_number(detail.get("quantity"))
            metrics["amount"] += parse_number(detail.get("discount_amount"))
    return summary


def summary_row(label: str, metrics: dict[str, object]) -> list[object]:
    order_count = len(metrics["order_numbers"])
    boxes = metrics["boxes"]
    amount = metrics["amount"]
    return [
        label,
        order_count,
        int(boxes),
        int(amount),
        round_half_up(amount, order_count),
        round_half_up(amount, boxes),
        round_half_up(boxes, order_count),
    ]


def write_summary(sheet, month_key: str, summary: dict[str, dict[str, object]]) -> int:
    title_row = 1
    sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=7)
    title_cell = sheet.cell(title_row, 1, f"{month_key[:4]}年{int(month_key[5:])}月訂單總覽")
    title_cell.fill = PatternFill("solid", fgColor="1F4E79")
    title_cell.font = Font(bold=True, color="FFFFFF", size=12)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[title_row].height = 25

    header_row = 2
    thin_black = Side(style="thin", color="000000")
    border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    for column, value in enumerate(SUMMARY_HEADERS, start=1):
        cell = sheet.cell(header_row, column, value)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[header_row].height = 24

    rows = [summary_row(category, summary[category]) for category in SUMMARY_CATEGORIES]
    total_metrics = empty_metrics()
    for category in SUMMARY_CATEGORIES:
        total_metrics["order_numbers"].update(summary[category]["order_numbers"])
        total_metrics["boxes"] += summary[category]["boxes"]
        total_metrics["amount"] += summary[category]["amount"]
    rows.append(summary_row("Total", total_metrics))

    for row_offset, values in enumerate(rows, start=1):
        row_number = header_row + row_offset
        label = str(values[0])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.fill = PatternFill("solid", fgColor=CATEGORY_COLORS[label])
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if column >= 2:
                cell.number_format = "#,##0"
            if label == "Total":
                cell.font = Font(bold=True)
        sheet.row_dimensions[row_number].height = 22

    return 9


def product_name_for(detail: dict, lookup: dict[str, str], unmapped: set[str]) -> str:
    product_code = str(detail.get("gift_code") or "").strip()
    product_name = lookup.get(product_code)
    if product_name:
        return product_name
    unmapped.add(product_code or "(空白代碼)")
    return f"未建檔（{product_code or '空白代碼'}）"


def add_month_sheet(
    workbook: Workbook,
    month_key: str,
    orders: list[dict],
    product_lookup: dict[str, str],
    unmapped_product_codes: set[str],
) -> int:
    sheet = workbook.create_sheet(month_key)
    sheet.sheet_view.showGridLines = False
    summary = calculate_month_summary(orders)
    detail_header_row = write_summary(sheet, month_key, summary)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E2EC")
    for column, value in enumerate(DETAIL_HEADERS, start=1):
        cell = sheet.cell(detail_header_row, column, value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin, right=thin)
    sheet.row_dimensions[detail_header_row].height = 26

    detail_count = 0
    row_number = detail_header_row + 1
    for order in orders:
        category = display_category(order.get("category"))
        store = STORE_BY_DEPARTMENT.get(
            str(order.get("department_code") or ""),
            str(order.get("department_code") or ""),
        )
        for detail in order.get("details") or []:
            values = [
                store,
                category,
                order.get("order_no") or "",
                parse_date(order.get("order_date")),
                product_name_for(detail, product_lookup, unmapped_product_codes),
                detail.get("list_price"),
                detail.get("quantity"),
                detail.get("total_amount"),
                detail.get("discount_amount"),
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, value)
                cell.fill = PatternFill("solid", fgColor=CATEGORY_COLORS[category])
                cell.border = Border(bottom=thin, right=thin)
                cell.alignment = Alignment(vertical="center")
                if column in (1, 2, 3, 4):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                if column >= 6:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            sheet.cell(row_number, 4).number_format = "yyyy-mm-dd"
            row_number += 1
            detail_count += 1

    widths = [12, 14, 18, 14, 42, 12, 10, 14, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = f"A{detail_header_row + 1}"
    sheet.auto_filter.ref = f"A{detail_header_row}:I{max(detail_header_row, row_number - 1)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return detail_count


def add_pending_sheet(workbook: Workbook, current_document: dict) -> int:
    sheet = workbook.create_sheet("待確認")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    headers = ["狀態", "圖片名稱", "訂單編號", "部門代號", "訂貨日期", "原因"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="BF9000")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    pending_entries = [
        entry
        for entry in current_document.get("orders") or []
        if entry.get("status") == "待確認"
    ]
    for entry in pending_entries:
        order = entry.get("order") or {}
        source = order.get("source") or {}
        images = entry.get("source_images") or []
        image_order_no = Path(str(images[0])).stem if images else ""
        sheet.append(
            [
                "待確認",
                ", ".join(str(image) for image in images),
                entry.get("order_no") or source.get("raw_order_no") or image_order_no,
                order.get("department_code") or "",
                order.get("order_date") or source.get("raw_order_date") or "",
                entry.get("reason") or source.get("error") or "OCR失敗或欄位不完整",
            ]
        )
    if not pending_entries:
        sheet.append(["完成", "", "", "", "", "目前沒有待確認資料"])

    sheet.auto_filter.ref = f"A1:F{max(1, sheet.max_row)}"
    for index, width in enumerate([12, 28, 20, 14, 14, 52], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    thin = Side(style="thin", color="D9E2EC")
    for row in sheet.iter_rows(min_row=2, max_col=6):
        for cell in row:
            cell.border = Border(bottom=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if pending_entries:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    return len(pending_entries)


def add_pending_sheet_v511(workbook: Workbook, current_document: dict, review_document: dict) -> int:
    if "待確認" in workbook.sheetnames:
        del workbook["待確認"]
    sheet = workbook.create_sheet("待確認")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    headers = [
        "狀態",
        "圖片名稱",
        "訂單編號",
        "部門代號",
        "訂貨日期",
        "原因",
        "處理結果",
        "處理備註",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="BF9000")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    records = {
        str(order_no): dict(record)
        for order_no, record in (review_document.get("records") or {}).items()
        if record.get("state") != "resolved"
    }
    for entry in current_document.get("orders") or []:
        if entry.get("status") != "待確認":
            continue
        order = entry.get("order") or {}
        source = order.get("source") or {}
        order_no = str(entry.get("order_no") or order.get("order_no") or source.get("raw_order_no") or "")
        if not order_no:
            continue
        record = records.setdefault(order_no, {})
        record.setdefault("order_no", order_no)
        record.setdefault("department_code", order.get("department_code") or "")
        record.setdefault("order_date", order.get("order_date") or source.get("raw_order_date") or "")
        record.setdefault("reason", entry.get("reason") or source.get("error") or "OCR 待確認")
        record.setdefault("source_images", entry.get("source_images") or [])
        record.setdefault("processing_result", "待確認")
        record.setdefault("processing_note", "")
        record.setdefault("state", "pending")

    ordered_records = sorted(records.values(), key=lambda record: str(record.get("order_no") or ""))
    for record in ordered_records:
        state = record.get("state") or "pending"
        status = "已確認不納入" if state == "excluded" else "待確認"
        processing_result = record.get("processing_result") or "待確認"
        if processing_result not in {"待確認", "重新辨識", "確認不納入"}:
            processing_result = "待確認"
        sheet.append(
            [
                status,
                ", ".join(str(value) for value in (record.get("source_images") or [])),
                record.get("order_no") or "",
                record.get("department_code") or "",
                record.get("order_date") or "",
                record.get("reason") or "",
                processing_result,
                record.get("processing_note") or "",
            ]
        )

    if not ordered_records:
        sheet.append(["無資料", "", "", "", "", "目前沒有待確認資料", "待確認", ""])

    last_row = max(2, sheet.max_row)
    validation = DataValidation(
        type="list",
        formula1='"待確認,重新辨識,確認不納入"',
        allow_blank=False,
    )
    validation.error = "請從下拉選單選擇處理結果。"
    validation.errorTitle = "處理結果不正確"
    validation.prompt = "選擇待確認、重新辨識或確認不納入。"
    validation.promptTitle = "人工處理"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"G2:G{last_row}")

    sheet.auto_filter.ref = f"A1:H{last_row}"
    widths = [16, 28, 20, 14, 14, 48, 18, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    thin = Side(style="thin", color="D9E2EC")
    for row in sheet.iter_rows(min_row=2, max_col=8):
        excluded = row[0].value == "已確認不納入"
        fill = PatternFill("solid", fgColor="E7E6E6" if excluded else "FFF2CC")
        for cell in row:
            cell.border = Border(bottom=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = fill
    return len(ordered_records)


def main() -> int:
    if len(sys.argv) < 3:
        raise SystemExit(
            "Usage: python build_current_orders_excel.py order_history.json current_orders.xlsx "
            "[current_orders.json] [preview.png]"
        )

    history_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    current_path = Path(sys.argv[3]) if len(sys.argv) >= 4 and sys.argv[3] else None
    history = json.loads(history_path.read_text(encoding="utf-8-sig"))
    current_document = (
        json.loads(current_path.read_text(encoding="utf-8-sig"))
        if current_path and current_path.exists()
        else {"orders": []}
    )
    review_path = history_path.parent / "pending_review.json"
    review_document = (
        json.loads(review_path.read_text(encoding="utf-8-sig"))
        if review_path.exists()
        else {"records": {}}
    )
    target_department = str(history.get("target_department") or "4001")
    product_lookup, lookup_status = load_or_refresh(MENUITEM_PATH, PRODUCT_LOOKUP_PATH)
    unmapped_product_codes: set[str] = set()

    orders = []
    for entry in (history.get("orders") or {}).values():
        order = entry.get("current_version") if isinstance(entry, dict) else None
        if not order:
            continue
        if str(order.get("department_code") or "") != target_department:
            continue
        if order.get("excel_eligible") is not True or order.get("voided") is True:
            continue
        orders.append(order)
    orders.sort(key=lambda order: (str(order.get("order_date") or ""), str(order.get("order_no") or "")))

    month_keys = recent_month_keys(report_as_of())
    orders_by_month: dict[str, list[dict]] = {month_key: [] for month_key in month_keys}
    for order in orders:
        order_date = parse_date(order.get("order_date"))
        if not order_date:
            continue
        month_key = order_date.strftime("%Y-%m")
        if month_key in orders_by_month:
            orders_by_month[month_key].append(order)
    populated_months = [month_key for month_key in month_keys if orders_by_month[month_key]]

    workbook = Workbook()
    workbook.remove(workbook.active)
    detail_row_count = 0
    for month_key in populated_months:
        detail_row_count += add_month_sheet(
            workbook,
            month_key,
            orders_by_month[month_key],
            product_lookup,
            unmapped_product_codes,
        )
    pending_count = add_pending_sheet_v511(workbook, current_document, review_document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(
        json.dumps(
            {
                "output": str(output_path),
                "targetDepartment": target_department,
                "store": STORE_BY_DEPARTMENT.get(target_department, target_department),
                "monthSheets": populated_months,
                "orderCount": sum(len(orders_by_month[month]) for month in populated_months),
                "detailRowCount": detail_row_count,
                "pendingCount": pending_count,
                "productLookupCount": len(product_lookup),
                "productLookupStatus": lookup_status,
                "unmappedProductCodes": sorted(unmapped_product_codes),
                "builder": "python-openpyxl",
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
