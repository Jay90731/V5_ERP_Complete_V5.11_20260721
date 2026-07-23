from __future__ import annotations

import argparse
import ctypes
import importlib.util
import json
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


VALID_RESULTS = {"待確認", "重新辨識", "確認不納入"}


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json_atomic(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)


def empty_review() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "updated_at": now_iso(),
        "records": {},
    }


def normalize_order_no(value: Any) -> str:
    return "".join(str(value or "").upper().split())


def image_order_no(raw: dict[str, Any]) -> str:
    image_name = str(raw.get("imageName") or "")
    stem = Path(image_name).stem.upper()
    return stem if stem.startswith("GOU") else normalize_order_no(raw.get("orderNo"))


def is_excel_locked(path: Path) -> bool:
    if not path.exists():
        return False
    if os.name != "nt":
        try:
            with path.open("r+b"):
                return False
        except PermissionError:
            return True

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        ctypes.c_wchar_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
    ]
    create_file.restype = ctypes.c_void_p
    handle = create_file(
        str(path),
        0x80000000 | 0x40000000,
        0,
        None,
        3,
        0x80,
        None,
    )
    invalid_handle = ctypes.c_void_p(-1).value
    if handle == invalid_handle:
        error = ctypes.get_last_error()
        if error in {32, 33}:
            return True
        raise OSError(error, f"無法檢查 Excel 狀態：{path}")
    kernel32.CloseHandle(ctypes.c_void_p(handle))
    return False


def assert_excel_closed(path: Path) -> None:
    if is_excel_locked(path):
        raise RuntimeError(
            "Excel 仍在開啟中。請先儲存並關閉 current_orders.xlsx，然後再按一次開始或繼續。"
        )


def record_from_current(entry: dict[str, Any], existing: dict[str, Any] | None = None) -> dict[str, Any]:
    order = entry.get("order") or {}
    source = order.get("source") or {}
    order_no = normalize_order_no(entry.get("order_no") or order.get("order_no") or source.get("raw_order_no"))
    current = dict(existing or {})
    timestamp = now_iso()
    current.update(
        {
            "order_no": order_no,
            "department_code": str(order.get("department_code") or current.get("department_code") or ""),
            "order_date": str(order.get("order_date") or source.get("raw_order_date") or current.get("order_date") or ""),
            "reason": str(entry.get("reason") or source.get("error") or current.get("reason") or ""),
            "source_images": list(entry.get("source_images") or current.get("source_images") or []),
            "source_paths": list(
                dict.fromkeys(
                    [
                        value
                        for value in [source.get("source_path"), *(current.get("source_paths") or [])]
                        if value
                    ]
                )
            ),
            "processing_result": current.get("processing_result") or "待確認",
            "processing_note": current.get("processing_note") or "",
            "state": current.get("state") if current.get("state") == "excluded" else "pending",
            "created_at": current.get("created_at") or timestamp,
            "updated_at": timestamp,
        }
    )
    return current


def merge_current_pending(review: dict[str, Any], current_document: dict[str, Any]) -> int:
    records = review.setdefault("records", {})
    added = 0
    for entry in current_document.get("orders") or []:
        if entry.get("status") != "待確認":
            continue
        order = entry.get("order") or {}
        order_no = normalize_order_no(entry.get("order_no") or order.get("order_no"))
        if not order_no:
            continue
        if order_no not in records:
            added += 1
        records[order_no] = record_from_current(entry, records.get(order_no))
    return added


def read_excel_marks(excel_path: Path) -> list[dict[str, str]]:
    if not excel_path.exists():
        return []
    workbook = load_workbook(excel_path, data_only=False, read_only=True)
    try:
        if "待確認" not in workbook.sheetnames:
            return []
        sheet = workbook["待確認"]
        headers = {
            str(cell.value or "").strip(): index
            for index, cell in enumerate(sheet[1], start=1)
        }
        order_col = headers.get("訂單編號")
        result_col = headers.get("處理結果")
        note_col = headers.get("處理備註")
        if not order_col:
            return []
        marks = []
        for row in range(2, sheet.max_row + 1):
            order_no = normalize_order_no(sheet.cell(row, order_col).value)
            if not order_no.startswith("GOU"):
                continue
            result = str(sheet.cell(row, result_col).value or "待確認").strip() if result_col else "待確認"
            note = str(sheet.cell(row, note_col).value or "").strip() if note_col else ""
            if result not in VALID_RESULTS:
                result = "待確認"
            marks.append({"order_no": order_no, "processing_result": result, "processing_note": note})
        return marks
    finally:
        workbook.close()


def command_sync(args: argparse.Namespace) -> dict[str, Any]:
    excel_path = Path(args.excel)
    review_path = Path(args.review_json)
    current_path = Path(args.current_orders)
    assert_excel_closed(excel_path)
    review = read_json(review_path, empty_review())
    review.setdefault("records", {})
    current = read_json(current_path, {"orders": []})
    added = merge_current_pending(review, current)
    marks = read_excel_marks(excel_path)
    timestamp = now_iso()
    for mark in marks:
        record = review["records"].setdefault(
            mark["order_no"],
            {
                "order_no": mark["order_no"],
                "created_at": timestamp,
                "state": "pending",
                "source_images": [],
                "source_paths": [],
            },
        )
        record["processing_result"] = mark["processing_result"]
        record["processing_note"] = mark["processing_note"]
        if mark["processing_result"] == "確認不納入":
            record["state"] = "excluded"
        elif record.get("state") != "resolved":
            record["state"] = "pending"
        record["updated_at"] = timestamp
    review["updated_at"] = timestamp
    write_json_atomic(review_path, review)
    return {
        "records": len(review["records"]),
        "new_pending": added,
        "excel_marks": len(marks),
        "retry_requested": sum(
            record.get("processing_result") == "重新辨識"
            and record.get("state") != "resolved"
            for record in review["records"].values()
        ),
        "excluded": sum(record.get("processing_result") == "確認不納入" for record in review["records"].values()),
    }


def load_extractor(path: Path):
    spec = importlib.util.spec_from_file_location("v5_pending_extractor", path)
    if not spec or not spec.loader:
        raise RuntimeError(f"無法載入 OCR 模組：{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def find_source_image(
    order_no: str,
    record: dict[str, Any],
    raw_by_order: dict[str, dict[str, Any]],
    current_by_order: dict[str, dict[str, Any]],
    runs_root: Path,
) -> Path | None:
    candidates: list[Path] = []
    for value in record.get("source_paths") or []:
        candidates.append(Path(value))
    raw = raw_by_order.get(order_no) or {}
    if raw.get("sourcePath"):
        candidates.append(Path(str(raw["sourcePath"])))
    current = current_by_order.get(order_no) or {}
    source = (current.get("order") or {}).get("source") or {}
    if source.get("source_path"):
        candidates.append(Path(str(source["source_path"])))
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    matches = sorted(
        runs_root.rglob(f"{order_no}.png"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return matches[0] if matches else None


def retry_order(extractor, image_path: Path, tesseract: Path, tessdata: Path) -> dict[str, Any]:
    order = extractor.extract_image(image_path, tesseract, tessdata)
    signature = extractor.image_signature(image_path)
    order["sourcePath"] = str(image_path)
    order["imageName"] = image_path.name
    order["extractorVersion"] = extractor.EXTRACTOR_VERSION
    order.update(signature)
    order["voided"] = extractor.image_has_void_stamp(image_path)
    order.pop("included", None)
    order.pop("statusLabel", None)
    order.pop("statusLevel", None)
    order.pop("errorReason", None)
    return extractor.validate_order(order)


def command_apply(args: argparse.Namespace) -> dict[str, Any]:
    review = read_json(Path(args.review_json), empty_review())
    records = review.setdefault("records", {})
    ocr_path = Path(args.ocr_json)
    payload = read_json(ocr_path, {"orders": []})
    current = read_json(Path(args.current_orders), {"orders": []})
    raw_by_order = {image_order_no(raw): raw for raw in payload.get("orders") or []}
    current_by_order = {
        normalize_order_no(entry.get("order_no")): entry
        for entry in current.get("orders") or []
        if normalize_order_no(entry.get("order_no"))
    }
    extractor = load_extractor(Path(args.extractor))
    timestamp = now_iso()
    report = {
        "retry_requested": 0,
        "retry_success": 0,
        "retry_pending": 0,
        "excluded": 0,
        "removed_from_batch": 0,
        "added_to_batch": 0,
        "replaced_in_batch": 0,
    }

    for order_no, record in records.items():
        result = record.get("processing_result") or "待確認"
        if result == "確認不納入":
            if order_no in raw_by_order:
                raw_by_order.pop(order_no, None)
                report["removed_from_batch"] += 1
            record["state"] = "excluded"
            record["last_result"] = "excluded"
            record["updated_at"] = timestamp
            report["excluded"] += 1
            continue
        if result != "重新辨識" or record.get("state") == "resolved":
            continue
        report["retry_requested"] += 1
        image_path = find_source_image(
            order_no,
            record,
            raw_by_order,
            current_by_order,
            Path(args.runs_root),
        )
        record["last_retry_at"] = timestamp
        if not image_path:
            record["state"] = "pending"
            record["last_retry_result"] = "image_not_found"
            record["reason"] = "找不到可重新辨識的原始圖片"
            record["updated_at"] = timestamp
            report["retry_pending"] += 1
            continue
        retried = retry_order(
            extractor,
            image_path,
            Path(args.tesseract),
            Path(args.tessdata),
        )
        existed = order_no in raw_by_order
        raw_by_order[order_no] = retried
        record["source_paths"] = list(
            dict.fromkeys([str(image_path), *(record.get("source_paths") or [])])
        )
        record["source_images"] = list(
            dict.fromkeys([image_path.name, *(record.get("source_images") or [])])
        )
        record["last_retry_result"] = retried.get("statusLevel") or "unknown"
        record["reason"] = retried.get("errorReason") or ""
        record["updated_at"] = timestamp
        if retried.get("statusLevel") == "success":
            record["state"] = "resolved"
            record["resolved_at"] = timestamp
            report["retry_success"] += 1
        else:
            record["state"] = "pending"
            report["retry_pending"] += 1
        if existed:
            report["replaced_in_batch"] += 1
        else:
            report["added_to_batch"] += 1

    final_orders = list(raw_by_order.values())
    for raw in final_orders:
        order_no = image_order_no(raw)
        record = records.get(order_no)
        if raw.get("statusLevel") == "success":
            if record and record.get("processing_result") != "確認不納入":
                record["state"] = "resolved"
                record["resolved_at"] = record.get("resolved_at") or timestamp
                record["updated_at"] = timestamp
            continue
        if not order_no:
            continue
        current_entry = {
            "order_no": order_no,
            "reason": raw.get("errorReason") or "OCR 待確認",
            "source_images": [raw.get("imageName")] if raw.get("imageName") else [],
            "order": {
                "order_no": order_no,
                "department_code": raw.get("departmentCode") or "",
                "order_date": raw.get("orderDate") or "",
                "source": {
                    "source_path": raw.get("sourcePath") or "",
                    "raw_order_no": raw.get("orderNo") or "",
                    "raw_order_date": raw.get("orderDate") or "",
                    "error": raw.get("errorReason") or "",
                },
            },
        }
        records[order_no] = record_from_current(current_entry, records.get(order_no))

    payload["orders"] = final_orders
    payload["pendingReviewAppliedAt"] = timestamp
    payload["pendingReviewReport"] = report
    write_json_atomic(ocr_path, payload)
    review["updated_at"] = timestamp
    write_json_atomic(Path(args.output_review), review)
    write_json_atomic(Path(args.report), report)
    return report


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="V5.11 待確認人工處理")
    subparsers = parser.add_subparsers(dest="command", required=True)

    lock = subparsers.add_parser("check-lock")
    lock.add_argument("--excel", required=True)

    sync = subparsers.add_parser("sync")
    sync.add_argument("--excel", required=True)
    sync.add_argument("--review-json", required=True)
    sync.add_argument("--current-orders", required=True)

    apply = subparsers.add_parser("apply")
    apply.add_argument("--review-json", required=True)
    apply.add_argument("--output-review", required=True)
    apply.add_argument("--current-orders", required=True)
    apply.add_argument("--ocr-json", required=True)
    apply.add_argument("--runs-root", required=True)
    apply.add_argument("--extractor", required=True)
    apply.add_argument("--tesseract", required=True)
    apply.add_argument("--tessdata", required=True)
    apply.add_argument("--report", required=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if args.command == "check-lock":
            assert_excel_closed(Path(args.excel))
            result = {"excel_closed": True}
        elif args.command == "sync":
            result = command_sync(args)
        else:
            result = command_apply(args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"待確認人工處理失敗：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
